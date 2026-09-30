from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..common.models import CanonicalObservation
from ..metadata.manifest import RawManifest

INDONESIAN_MONTHS = {
    "januari": 1,
    "februari": 2,
    "maret": 3,
    "april": 4,
    "mei": 5,
    "juni": 6,
    "juli": 7,
    "agustus": 8,
    "september": 9,
    "oktober": 10,
    "november": 11,
    "desember": 12,
}
EXPECTED_HEADERS = ("NO", "Tanggal", "BI-7Day-RR")


def _decision_date(value: object, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        raise TypeError(f"row {row_number}: invalid decision date {value!r}")
    parts = value.replace("\u00a0", " ").strip().split()
    if len(parts) != 3:
        raise ValueError(f"row {row_number}: invalid Indonesian date {value!r}")
    month = INDONESIAN_MONTHS.get(parts[1].casefold())
    if month is None:
        raise ValueError(f"row {row_number}: unknown Indonesian month {parts[1]!r}")
    try:
        return date(int(parts[2]), month, int(parts[0]))
    except ValueError as exc:
        raise ValueError(f"row {row_number}: invalid Indonesian date {value!r}") from exc


def _percent(value: object, row_number: int) -> Decimal:
    if not isinstance(value, str) or not value.strip().endswith("%"):
        raise ValueError(f"row {row_number}: expected a percentage string, got {value!r}")
    normalized = value.strip()[:-1].strip().replace(",", ".")
    try:
        return Decimal(normalized)
    except InvalidOperation as exc:
        raise ValueError(f"row {row_number}: invalid policy rate {value!r}") from exc


def transform_bi_policy_rate_xlsx(
    path: Path,
    *,
    source: dict[str, Any],
    manifest: RawManifest,
) -> list[CanonicalObservation]:
    """Transform the documented BI-7Day-RR workbook layout into event observations."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "BI-7Day-RR" not in workbook.sheetnames:
            raise ValueError("BI workbook is missing the expected 'BI-7Day-RR' worksheet")
        sheet = workbook["BI-7Day-RR"]
        headers = tuple(sheet.cell(row=5, column=column).value for column in range(1, 4))
        if headers != EXPECTED_HEADERS:
            raise ValueError(f"unexpected BI workbook headers at row 5: {headers!r}")

        observations: list[CanonicalObservation] = []
        previous_date: date | None = None
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=6, min_col=1, max_col=3, values_only=True), start=6
        ):
            if all(value is None for value in values):
                continue
            sequence, date_value, rate_value = values
            expected_sequence = len(observations) + 1
            if sequence != expected_sequence:
                raise ValueError(
                    f"row {row_number}: expected sequence {expected_sequence}, got {sequence!r}"
                )
            decision_date = _decision_date(date_value, row_number)
            if previous_date is not None and decision_date >= previous_date:
                raise ValueError(
                    f"row {row_number}: decision dates must be unique and newest-to-oldest"
                )
            previous_date = decision_date
            observations.append(
                CanonicalObservation(
                    indicator_code=source["indicator_code"],
                    indicator_name=source["indicator_name"],
                    geography_code=source["geography_code"],
                    geography_name=source["geography_name"],
                    geography_level=source["geography_level"],
                    period_start=decision_date,
                    period_end=decision_date,
                    frequency=source["frequency"],
                    value=_percent(rate_value, row_number),
                    unit_code=source["unit_code"],
                    unit_name=source["unit_name"],
                    institution=source["institution"],
                    dataset_code=source["dataset_code"],
                    dataset_name=source["dataset_name"],
                    source_url=source["source_url"],
                    access_method=source["access_method"],
                    retrieved_at=datetime.fromisoformat(manifest.retrieved_at),
                    raw_file_path=manifest.raw_file_path,
                    checksum_sha256=manifest.checksum_sha256,
                    http_status=manifest.http_status,
                )
            )
        if not observations:
            raise ValueError("BI workbook contains no policy-rate observations")
        return observations
    finally:
        workbook.close()
